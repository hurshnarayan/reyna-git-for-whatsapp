from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


SOURCE = Path("/Users/harsh/Documents/Reyna_SIH260150.pptx")
OUTPUT = Path("/Users/harsh/code/reyna/Reyna_SIH260150_updated.pptx")

NAVY = RGBColor(22, 50, 79)
BLUE = RGBColor(61, 90, 254)
GREEN = RGBColor(0, 155, 102)
ORANGE = RGBColor(245, 124, 0)
INK = RGBColor(28, 35, 43)
MUTED = RGBColor(87, 101, 112)
WHITE = RGBColor(255, 255, 255)
PALE_BLUE = RGBColor(239, 243, 255)
PALE_GREEN = RGBColor(235, 248, 242)
PALE_ORANGE = RGBColor(255, 244, 229)
PALE_GREY = RGBColor(244, 246, 248)


def remove(shape):
    element = shape._element
    element.getparent().remove(element)


def set_fill(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def set_line(shape, color, width=1):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def text(shape, value, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         valign=MSO_ANCHOR.TOP, font="Aptos", margins=(0.04, 0.04, 0.04, 0.04),
         spacing=1.04):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margins[0])
    tf.margin_top = Inches(margins[1])
    tf.margin_right = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    tf.vertical_anchor = valign
    lines = value.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(max(0, size * (spacing - 1)))
        p.line_spacing = spacing
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color


def card(shape, fill=WHITE, line=RGBColor(222, 228, 233)):
    set_fill(shape, fill)
    set_line(shape, line, 1)


def add_pill(slide, x, y, w, h, label, fill, color=WHITE, size=11):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    set_fill(sh, fill)
    sh.line.fill.background()
    text(sh, label, size=size, color=color, bold=True, align=PP_ALIGN.CENTER,
         valign=MSO_ANCHOR.MIDDLE, margins=(0.06, 0.01, 0.06, 0.01))
    return sh


def style_heading(slide, title, subtitle):
    shapes = list(slide.shapes)
    text(shapes[6], title, 23, NAVY, True, valign=MSO_ANCHOR.MIDDLE)
    text(shapes[7], subtitle, 10.5, MUTED, False, valign=MSO_ANCHOR.MIDDLE)


prs = Presentation(SOURCE)
prs.core_properties.title = "Retina App: Architecture & Working Overview"
prs.core_properties.subject = "SIH260150 technical architecture presentation"

# Slide 1 — title
s = prs.slides[0]
sh = list(s.shapes)
text(sh[4], "Retina", 45, NAVY, True, valign=MSO_ANCHOR.MIDDLE)
text(sh[5], "Your WhatsApp media, searchable and conversational.", 18, MUTED,
     valign=MSO_ANCHOR.MIDDLE)
text(sh[18], "Reyna", 15, INK, True, valign=MSO_ANCHOR.MIDDLE)
text(sh[21], "TEAM REYNA", 11, NAVY, True, valign=MSO_ANCHOR.MIDDLE)

# Slide 2 — core concept and ingestion
s = prs.slides[1]
sh = list(s.shapes)
style_heading(
    s,
    "RETINA · THE SECOND BRAIN FOR WHATSAPP MEDIA",
    "Core concept  ·  background ingestion  ·  sender-aware context",
)
text(sh[9], "Incoming images and PDFs become searchable knowledge—without a manual upload step.",
     17, NAVY, True, valign=MSO_ANCHOR.MIDDLE)

card(sh[10], PALE_BLUE, RGBColor(207, 216, 255))
text(sh[11], "BACKGROUND LISTENER", 11, BLUE, True)
text(
    sh[12],
    "• WhatsAppListenerService.kt runs as an Android NotificationListenerService.\n"
    "• It monitors incoming WhatsApp notifications in the background.\n"
    "• It targets notifications that carry a Photo or Document attachment.",
    14, INK, spacing=1.12,
)

card(sh[13], PALE_GREEN, RGBColor(196, 231, 213))
text(sh[14], "SMART SENDER PARSING", 11, GREEN, True)
text(
    sh[15],
    "• Direct chats: captures the sender.\n"
    "• Group chats: captures both the group and the individual sender.\n"
    "• Example: “Nidhi in SIH'26 : Selected teams”.",
    14, INK, spacing=1.12,
)

card(sh[16], PALE_ORANGE, RGBColor(255, 215, 166))
text(sh[17], "ZERO-TOUCH KNOWLEDGE PIPELINE", 11, ORANGE, True)
remove(sh[18])
text(
    sh[19],
    "The listener triggers extraction, AI understanding, local storage and retrieval—silently in the background.",
    13.5, INK, valign=MSO_ANCHOR.BOTTOM,
)
for i, (label, fill) in enumerate([
    ("1  Detect", BLUE), ("2  Read", GREEN), ("3  Embed", ORANGE), ("4  Retrieve", NAVY)
]):
    add_pill(s, 8.78, 2.82 + i * 0.66, 3.62, 0.48, label, fill, size=12)

# Slide 3 — technical pipeline
s = prs.slides[2]
sh = list(s.shapes)
style_heading(
    s,
    "TECHNICAL APPROACH",
    "Android ingestion  ·  OCR and parsing  ·  Gemini embeddings  ·  local retrieval",
)
steps = [
    ("1", "Listen\nWhatsApp\nnotification"),
    ("2", "Identify\nPhoto or\nDocument"),
    ("3", "Extract\nML Kit or\nPdfRenderer"),
    ("4", "Embed\nGemini\nVectorEngine"),
    ("5", "Store\nLocal JSON\ndatabase"),
    ("6", "Ask\nRetrieve and\nanswer"),
]
for j, (number, label) in enumerate(steps):
    bg, nbox, lbox = sh[9 + j * 3], sh[10 + j * 3], sh[11 + j * 3]
    card(bg, PALE_BLUE if j < 2 else PALE_GREEN if j < 4 else PALE_ORANGE,
         RGBColor(214, 220, 230))
    text(nbox, number, 12, BLUE if j < 2 else GREEN if j < 4 else ORANGE, True)
    text(lbox, label, 12.5, INK, True, valign=MSO_ANCHOR.MIDDLE)

card(sh[27], PALE_BLUE, RGBColor(207, 216, 255))
text(sh[28], "EXTRACT", 11, BLUE, True)
text(
    sh[29],
    "Images\nGoogle ML Kit OCR scans the image and returns readable text.\n\n"
    "PDFs / documents\nAndroid PdfRenderer processes document pages for text extraction.",
    13.2, INK, spacing=1.08,
)

card(sh[30], PALE_GREEN, RGBColor(196, 231, 213))
text(sh[31], "UNDERSTAND & STORE", 11, GREEN, True)
text(
    sh[32],
    "VectorEngine.kt bundles extracted text + filename + sender.\n\n"
    "Gemini converts that context into a semantic embedding.\n\n"
    "FileStorage.kt saves metadata, text and vectors in a local JSON database.",
    13.2, INK, spacing=1.08,
)
sh[32].top = Inches(3.3)
sh[32].height = Inches(3.15)
for idx in range(33, 40):
    remove(sh[idx])

card(sh[40], PALE_ORANGE, RGBColor(255, 215, 166))
text(sh[41], "SEARCH & ANSWER", 11, ORANGE, True)
remove(sh[42])
text(
    sh[43],
    "The query becomes an embedding. Retina ranks every file using cosine similarity, then boosts explicit sender matches.\n\n"
    "The top 3 files, question and chat history go to Gemini Chat for a grounded response.",
    13.5, INK, spacing=1.1,
)
sh[43].top = Inches(3.3)
sh[43].height = Inches(3.15)

# Slide 4 — persistence and concurrency
s = prs.slides[3]
sh = list(s.shapes)
style_heading(
    s,
    "PERSISTENCE & RELIABILITY",
    "Local file memory  ·  thread-safe writes  ·  durable chat history",
)
cards = [
    (9, 10, 11, "FILE DATABASE", PALE_BLUE, BLUE,
     "FileStorage.kt keeps each file's metadata, extracted text and embedding in a local JSON database on the device."),
    (12, 13, 14, "SAFE BACKGROUND WRITES", PALE_GREEN, GREEN,
     "WhatsApp files can arrive rapidly. @Synchronized locks serialize database access and protect the JSON file from overlapping writes and corruption."),
    (15, 16, 17, "CONVERSATION MEMORY", PALE_ORANGE, ORANGE,
     "ChatStorage.kt saves conversations locally. When the app restarts, earlier interactions remain available as context for the next answer."),
]
for bg_i, title_i, body_i, title_value, fill, accent, body in cards:
    card(sh[bg_i], fill, RGBColor(214, 220, 230))
    text(sh[title_i], title_value, 11, accent, True)
    text(sh[body_i], body, 15, INK, spacing=1.15)
card(sh[18], NAVY, NAVY)
text(
    sh[19],
    "Local metadata + extracted text + embeddings + chat history = a persistent second brain on the device.",
    15, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
)

# Slide 5 — hybrid search and conversational behavior
s = prs.slides[4]
sh = list(s.shapes)
style_heading(
    s,
    "INTELLIGENT SEARCH & CHAT",
    "Semantic retrieval  ·  sender-aware ranking  ·  fetch-aware interface",
)
metrics = [
    (10, 11, "MEANING", "query embedding + cosine similarity", BLUE, PALE_BLUE),
    (13, 14, "SENDER", "keyword boost for by/from [Name]", GREEN, PALE_GREEN),
    (16, 17, "TOP 3", "highest-ranked files become context", ORANGE, PALE_ORANGE),
    (19, 20, "FETCH", "shows matching files as chips", NAVY, PALE_GREY),
]
for idx, (value_i, label_i, value, label, accent, fill) in enumerate(metrics):
    bg_i = 9 + idx * 3
    card(sh[bg_i], fill, RGBColor(214, 220, 230))
    text(sh[value_i], value, 19, accent, True, valign=MSO_ANCHOR.MIDDLE)
    text(sh[label_i], label, 10.5, MUTED, valign=MSO_ANCHOR.MIDDLE)

card(sh[21], PALE_BLUE, RGBColor(207, 216, 255))
text(sh[22], "HYBRID RETRIEVAL", 11, BLUE, True)
text(
    sh[23],
    "1. Embed the user's query.\n"
    "2. Compare it with every stored file using cosine similarity.\n"
    "3. Detect explicit sender phrases such as “by Nidhi” or “from Nidhi”.\n"
    "4. Apply a strong sender boost and rank the results.",
    13.5, INK, spacing=1.08,
)

card(sh[24], PALE_GREEN, RGBColor(196, 231, 213))
text(sh[25], "CONTEXTUAL ANSWERING", 11, GREEN, True)
text(
    sh[26],
    "• With “fetch”: surface the top 3 files as interactive chips.\n"
    "• Without “fetch”: retrieve silently and keep the chat uncluttered.\n"
    "• Send the top 3 files + question + chat history to Gemini Chat.\n"
    "• Generate an answer grounded in the exact file contents.",
    13.5, INK, spacing=1.08,
)
card(sh[27], NAVY, NAVY)
text(
    sh[28],
    "Example: “Fetch the image from Nidhi about isometric projection.”",
    14.5, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
)

# Slide 6 — architecture map and component ownership
s = prs.slides[5]
sh = list(s.shapes)
style_heading(
    s,
    "END-TO-END ARCHITECTURE",
    "Named components and the data passed between them",
)
card(sh[9], PALE_BLUE, RGBColor(207, 216, 255))
text(sh[10], "ANDROID INGESTION & EXTRACTION", 11, BLUE, True)
text(
    sh[11],
    "WhatsAppListenerService.kt\nNotification interception + attachment and sender parsing\n\n"
    "Google ML Kit OCR\nText extraction from images\n\n"
    "Android PdfRenderer\nDocument-page processing",
    13.2, INK, spacing=1.05,
)

card(sh[12], PALE_GREEN, RGBColor(196, 231, 213))
text(sh[13], "AI, STORAGE & MEMORY", 11, GREEN, True)
text(
    sh[14],
    "VectorEngine.kt\nGemini embeddings for files and queries\n\n"
    "FileStorage.kt\nLocal JSON metadata, text and vectors\n\n"
    "ChatStorage.kt\nPersistent local conversation history",
    13.2, INK, spacing=1.05,
)

card(sh[15], NAVY, NAVY)
text(sh[16], "HOW ONE FILE BECOMES AN ANSWER", 11, RGBColor(164, 185, 255), True)
text(
    sh[17],
    "WhatsApp notification  →  attachment + sender  →  OCR/PDF text  →  Gemini embedding  →  local JSON  →  hybrid Top-3 retrieval  →  Gemini Chat answer",
    16, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
)

# Remove stale speaker notes from the previous architecture.
for slide in prs.slides:
    try:
        tf = slide.notes_slide.notes_text_frame
        if tf is not None:
            tf.text = ""
    except Exception:
        pass

prs.save(OUTPUT)
print(OUTPUT)
